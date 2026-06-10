# -*- coding: utf-8 -*-
"""
build_vocab.py — Pha 1 của app học tiếng Trung.

Đọc file Excel HSK1–6 (6 sheet) -> sinh data/words.json đã gắn sẵn:
  - struct_group  (S1–S5 theo số chữ)
  - semantic_group qua pipeline 2 lớp đầu:
        Lớp 1: danh sách từ chức năng đóng -> F1–F6
        Lớp 2: tra từ khóa trong nghĩa tiếng Việt -> A1–E2
  - traditional (OpenCC s2t), han_viet (nếu Excel có dạng "pinyin [hán việt]")

Các từ không khớp Lớp 1/2 -> semantic_group = null, chờ Lớp 3 (Qwen3) sau này.

Cách chạy (Windows PowerShell, trong venv):
    pip install openpyxl opencc-python-reimplemented
    python scripts/build_vocab.py "đường/dẫn/FULL_TU_VUNG_HSK1_HSK6.xlsx"

Nếu không truyền tham số, script tìm backend/data/vocab.xlsx hoặc data/vocab.xlsx.
"""

import sys
import re
import json
import os
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("Thiếu openpyxl. Chạy: pip install openpyxl")

try:
    from opencc import OpenCC
    _cc = OpenCC("s2t")
    def to_traditional(s):
        return _cc.convert(s)
except Exception:
    # Không có opencc -> bỏ qua phồn thể, không chặn pipeline
    def to_traditional(s):
        return None


# ---------------------------------------------------------------------------
# Làm sạch chữ Hán: bỏ ngoặc chú thích, giữ đúng dạng từ
# ---------------------------------------------------------------------------
# 5 ca đặc biệt có ngoặc trong file gốc (xử lý tường minh cho chắc chắn)
PAREN_OVERRIDES = {
    "（杯）子": "杯子",      # ngoặc bọc 1 chữ thuộc về từ -> giữ lại
    "哪里 (哪儿)": "哪里",   # ngoặc là biến thể -> bỏ
    "椅子 （把）": "椅子",   # ngoặc là gợi ý lượng từ -> bỏ
    "得（助动词）": "得",     # ngoặc là chú từ loại -> bỏ
    "等（助词）": "等",       # ngoặc là chú từ loại -> bỏ
}

def clean_word(raw):
    raw = str(raw).strip()
    if raw in PAREN_OVERRIDES:
        return PAREN_OVERRIDES[raw]
    # Quy tắc chung: bỏ phần trong ngoặc (cả nửa/đầy đủ) và khoảng trắng thừa
    out = re.sub(r"[（(][^）)]*[）)]", "", raw)
    out = re.sub(r"\s+", "", out)
    return out or raw


def split_pinyin_hanviet(raw):
    """ 'ài qíng [ái tình]' -> ('ài qíng', 'ái tình'); 'ài' -> ('ài', None) """
    if raw is None:
        return None, None
    raw = str(raw).strip()
    m = re.search(r"\[([^\]]+)\]", raw)
    if m:
        han_viet = m.group(1).strip()
        pinyin = raw[:m.start()].strip()
    else:
        han_viet = None
        pinyin = raw
    # bỏ chú thích trong ngoặc còn sót ở cột phiên âm (vd "de （助动词）", "nǎlǐ (nǎr)")
    pinyin = re.sub(r"[（(][^）)]*[）)]", "", pinyin)
    pinyin = re.sub(r"\s+", " ", pinyin).strip()
    return (pinyin or None), han_viet


# ---------------------------------------------------------------------------
# CẤP 1 — struct_group theo số chữ Hán
# ---------------------------------------------------------------------------
def struct_group(word):
    n = len(word)
    if n <= 1:
        return "S1"
    if n == 2:
        return "S2"
    if n == 3:
        return "S3"
    if n == 4:
        return "S4"   # 4 chữ -> mặc định coi là thành ngữ (tinh chỉnh bằng CC-CEDICT sau)
    return "S5"        # 5+ chữ -> cụm dài


# ---------------------------------------------------------------------------
# LỚP 1 — từ chức năng (danh sách đóng) -> F1–F6
# Nguồn: VOCAB_CLASSIFICATION_SYSTEM.md, Phần 2, Nhóm F
# ---------------------------------------------------------------------------
FUNCTION_WORDS = {
    "F1": ["和", "但是", "因为", "所以", "虽然", "如果", "既然", "否则", "而且", "不过"],
    "F2": ["在", "从", "到", "对", "把", "被", "为了", "关于", "按照", "除了"],
    "F3": ["了", "过", "着", "吗", "吧", "呢", "啊", "的", "地", "得", "嘛", "啦", "哟"],
    "F4": ["很", "非常", "太", "最", "也", "都", "就", "才", "还", "已经", "一直", "一定", "曾经", "突然"],
    "F5": ["我", "你", "他", "她", "它", "我们", "你们", "他们", "这", "那", "谁", "什么", "哪", "怎么"],
    "F6": ["个", "本", "张", "条", "件", "辆", "只", "把", "次", "遍", "种", "位", "名", "块", "点", "些"],
}
# Tra ngược: chữ -> mã F (ưu tiên F1..F6 theo thứ tự, first-match thắng)
FUNCTION_LOOKUP = {}
for code in ["F1", "F2", "F3", "F4", "F5", "F6"]:
    for w in FUNCTION_WORDS[code]:
        FUNCTION_LOOKUP.setdefault(w, code)


def layer1_function(word):
    return FUNCTION_LOOKUP.get(word)


# ---------------------------------------------------------------------------
# LỚP 2 — từ khóa trong nghĩa tiếng Việt -> A1–E2
# Nguồn: VOCAB_CLASSIFICATION_SYSTEM.md, Phần 3, bảng từ khóa
# ---------------------------------------------------------------------------
KEYWORD_MAP = [
    ("A1", ["bệnh", "ung thư", "đau", "sốt", "thuốc", "y tế", "phẫu thuật", "bác sĩ", "bệnh viện", "cơ thể", "sức khỏe"]),
    ("A2", ["vui", "buồn", "sợ", "tức", "yêu", "ghét", "lo", "cảm xúc", "tâm lý", "tính cách", "cảm giác", "tâm trạng"]),
    ("A3", ["gia đình", "bạn bè", "đồng nghiệp", "chào", "xin lỗi", "cảm ơn", "quan hệ", "giao tiếp", "lịch sự"]),
    ("A4", ["ăn", "mặc", "ngủ", "nghỉ", "mua", "giải trí", "hàng ngày", "sinh hoạt", "nhà ở", "đi lại"]),
    ("B1", ["học", "thi", "trường", "giáo viên", "môn học", "nghiên cứu", "kiến thức", "giáo dục", "sinh viên"]),
    ("B2", ["làm việc", "công ty", "nghề", "chức vụ", "hợp đồng", "nhiệm vụ", "nhân viên", "quản lý"]),
    ("B3", ["tiền", "ngân hàng", "kinh tế", "thị trường", "đầu tư", "lợi nhuận", "tiêu dùng", "thu nhập", "giá"]),
    ("C1", ["suy nghĩ", "phân tích", "lý luận", "quan điểm", "nguyên tắc", "lý thuyết", "phán đoán", "logic"]),
    ("C2", ["ngôn ngữ", "dịch", "chữ", "văn", "giải thích", "diễn đạt", "mô tả", "từ ngữ", "câu"]),
    ("D1", ["núi", "sông", "biển", "thời tiết", "động vật", "thực vật", "môi trường", "ô nhiễm", "tài nguyên"]),
    ("D2", ["trên", "dưới", "trái", "phải", "trước", "sau", "thời gian", "khi", "lúc", "thường xuyên", "đôi khi", "gần đây"]),
    ("E1", ["luật", "chính phủ", "quyền", "bầu cử", "chính sách", "dân chủ", "nhà nước", "pháp luật"]),
    ("E2", ["văn hóa", "lịch sử", "nghệ thuật", "âm nhạc", "hội họa", "lễ hội", "phong tục", "truyền thống"]),
]

def _kw_hit(keyword, tokens, text):
    """Từ khóa 1 tiếng -> khớp theo token (tránh khớp nhầm chuỗi con);
    từ khóa nhiều tiếng -> khớp chuỗi con."""
    if " " in keyword:
        return keyword in text
    return keyword in tokens

def layer2_keyword(meaning_vi):
    if not meaning_vi:
        return None
    text = meaning_vi.lower()
    # tách token theo ký tự không phải chữ cái (giữ dấu tiếng Việt)
    tokens = set(re.split(r"[^0-9a-zà-ỹ]+", text))
    for code, kws in KEYWORD_MAP:
        for kw in kws:
            if _kw_hit(kw, tokens, text):
                return code
    return None


# ---------------------------------------------------------------------------
# Đọc Excel + chạy pipeline
# ---------------------------------------------------------------------------
def find_excel(arg):
    if arg and os.path.isfile(arg):
        return arg
    for c in ["backend/data/vocab.xlsx", "data/vocab.xlsx", "vocab.xlsx"]:
        if os.path.isfile(c):
            return c
    sys.exit("Không tìm thấy file Excel. Truyền đường dẫn: python scripts/build_vocab.py <file.xlsx>")


def cell(row, i):
    return row[i] if i < len(row) else None


def main():
    src = find_excel(sys.argv[1] if len(sys.argv) > 1 else None)
    print(f"Đọc Excel: {src}")
    wb = openpyxl.load_workbook(src, read_only=True)

    words = []
    stats = {"S": {}, "sem": {}, "src": {"rule": 0, "keyword": 0, "unclassified": 0}}

    for sn in wb.sheetnames:
        m = re.search(r"(\d)", sn)
        if not m:
            continue
        level = int(m.group(1))
        ws = wb[sn]
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_word = cell(row, 1)
            if raw_word is None or str(raw_word).strip() == "":
                continue
            stt = cell(row, 0)
            try:
                stt = int(stt)
            except (TypeError, ValueError):
                stt = len(words) + 1

            simplified = clean_word(raw_word)
            pinyin, han_viet = split_pinyin_hanviet(cell(row, 2))
            meaning_vi = (str(cell(row, 3)).strip() if cell(row, 3) else None)
            example_zh = (str(cell(row, 4)).strip() if cell(row, 4) else None)
            example_pinyin = (str(cell(row, 5)).strip() if cell(row, 5) else None)
            example_vi = (str(cell(row, 6)).strip() if cell(row, 6) else None)

            sg = struct_group(simplified)

            # --- pipeline phân loại ---
            sem = layer1_function(simplified)
            if sem:
                csource, conf = "rule", 1.0
            else:
                sem = layer2_keyword(meaning_vi)
                if sem:
                    csource, conf = "keyword", 0.85
                else:
                    csource, conf = None, None

            if csource == "rule":
                stats["src"]["rule"] += 1
            elif csource == "keyword":
                stats["src"]["keyword"] += 1
            else:
                stats["src"]["unclassified"] += 1
            stats["S"][sg] = stats["S"].get(sg, 0) + 1
            if sem:
                stats["sem"][sem] = stats["sem"].get(sem, 0) + 1

            words.append({
                "id": f"h{level}-{stt:04d}",
                "simplified": simplified,
                "traditional": to_traditional(simplified),
                "pinyin": pinyin,
                "han_viet": han_viet,
                "meaning_vi": meaning_vi,
                "example_zh": example_zh,
                "example_pinyin": example_pinyin,
                "example_vi": example_vi,
                "hsk_level": level,
                "source": "excel",
                "struct_group": sg,
                "semantic_group": sem,
                "classification_source": csource,
                "classification_conf": conf,
                "needs_review": False,
                "context_tags": None,
            })

    out = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "count": len(words),
        "words": words,
    }
    out_path = os.path.join("data", "words.json")
    os.makedirs("data", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=0)

    # --- báo cáo ---
    total = len(words)
    print(f"\nĐã ghi {out_path} — {total} từ")
    print("\nstruct_group:", dict(sorted(stats["S"].items())))
    print("\nNguồn phân loại semantic:")
    for k in ("rule", "keyword", "unclassified"):
        v = stats["src"][k]
        print(f"  {k:13s}: {v:5d}  ({v*100//total}%)")
    print("\nsemantic_group đã gán:", dict(sorted(stats["sem"].items())))


if __name__ == "__main__":
    main()
